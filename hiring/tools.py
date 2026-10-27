"""Plain-code tools the AI employees call. No model is used in this file.

- read_cv: get text out of PDF, Word and text files
- rule_profile: a simple profile made by rules (used when the model fails)
- score_cv / skill gaps / best_area_by_rules: the weighted score
- pick_interview_slot: next free slot on a working day
- induction_checklist and the reply templates
"""
from __future__ import annotations

import hashlib
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from . import config
from .checks import EMAIL_RE, CVText, norm

# ---- Reading CV files -----------------------------------------------------------
def read_cv(path: Path) -> str:
    suffix = path.suffix.lower()
    text = ""
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix == ".txt":
        text = path.read_bytes().decode("utf-8", errors="replace")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[  ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# ---- A profile made by rules ---------------------------------------------------------
HEADINGS = {
    "education": "education", "academic qualifications": "education", "qualifications": "education",
    "skills": "skills", "technical skills": "skills", "key skills": "skills", "core skills": "skills",
    "projects": "projects", "academic projects": "projects", "personal projects": "projects",
    "experience": "experience", "work experience": "experience", "internships": "experience",
    "employment": "experience", "professional experience": "experience",
    "summary": "other", "profile": "other", "about me": "other", "certifications": "other",
    "languages": "other", "references": "other", "achievements": "other", "contact": "other",
    "interests": "other", "activities": "other", "extracurricular activities": "other",
}
GENERAL_SKILLS = ["Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "SQL", "Git", "Linux", "HTML",
                  "CSS", "React", "Node.js", "Excel", "Power BI", "Tableau", "Figma", "Communication",
                  "Recruitment", "Payroll", "Accounting", "Marketing", "Sales", "Networking"]


def split_sections(cv_text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"top": []}
    current = "top"
    for line in cv_text.split("\n"):
        key = norm(line.rstrip(":")).strip()
        if key in HEADINGS and len(line) < 40:
            current = HEADINGS[key]
            sections.setdefault(current, [])
            continue
        if line.strip():
            sections.setdefault(current, []).append(line.strip())
    return sections


def _items(lines: list[str]) -> list[str]:
    out = []
    for line in lines:
        item = re.sub(r"^[\-•*●▪>]+\s*", "", line).strip()
        if item and item not in out:
            out.append(item)
    return out


def skill_vocabulary(areas: list) -> list[str]:
    vocab = list(GENERAL_SKILLS)
    for a in areas:
        for s in list(a.required_skills or []) + list(a.preferred_skills or []):
            if s not in vocab:
                vocab.append(s)
    return vocab


def rule_profile(cv_text: str, vocab: list[str]) -> dict:
    sec = split_sections(cv_text)
    cv = CVText(cv_text)
    top = sec.get("top", [])
    name = ""
    for line in top[:5]:
        words = line.split()
        if 2 <= len(words) <= 5 and not EMAIL_RE.search(line) and not re.search(r"\d", line):
            name = line.strip()
            break
    m = EMAIL_RE.search(cv_text)
    phone = ""
    pm = re.search(r"(\+?\d[\d\s\-]{8,}\d)", cv_text)
    if pm:
        phone = pm.group(1).strip()
    skills = [s for s in vocab if cv.has(s)]
    for line in sec.get("skills", []):
        for part in re.split(r"[,;|/]", re.sub(r"^[^:]{0,30}:", "", line)):
            part = part.strip(" -•*.")
            if 1 < len(part) <= 40 and part not in skills:
                skills.append(part)
    return {
        "name": name, "email": m.group(0) if m else "", "phone": phone,
        "education": _items(sec.get("education", []))[:8],
        "skills": skills[:25],
        "projects": _items(sec.get("projects", []))[:10],
        "experience": _items(sec.get("experience", []))[:10],
    }


# ---- Scoring ---------------------------------------------------------------------------
def _area_terms(area) -> list[str]:
    return list(area.required_skills or []) + list(area.preferred_skills or []) + list(area.keywords or [])


def score_cv(profile: dict, cv_text: str, area) -> tuple[int | None, dict, list[str]]:
    """Weighted score out of 100 against one area. Returns (score, parts, gaps).
    The score is None when the area has no required skills yet."""
    w = config.WEIGHTS
    required = list(area.required_skills or [])
    if not required:
        return None, {}, []
    cv = CVText(cv_text)
    found = [s for s in required if cv.has(s)]
    gaps = [s for s in required if s not in found]
    preferred_found = [s for s in (area.preferred_skills or []) if cv.has(s)]
    skills_pts = w["skills"] * len(found) / len(required)

    terms = _area_terms(area)
    projects = profile.get("projects") or []
    relevant_projects = [p for p in projects if any(CVText(p).has(t) for t in terms)]
    projects_pts = w["projects"] * min(len(relevant_projects), 2) / 2

    education = " ".join(profile.get("education") or [])
    edu_norm = " " + norm(education) + " "
    if any(f" {norm(k)} " in edu_norm for k in config.RELEVANT_EDUCATION):
        edu_pts, edu_why = w["education"], "Studies a related subject"
    elif education.strip():
        edu_pts, edu_why = w["education"] / 2, "Studies, but not in a related subject"
    else:
        edu_pts, edu_why = 0, "No education found in the CV"

    experience = profile.get("experience") or []
    relevant_exp = [e for e in experience if any(CVText(e).has(t) for t in terms)]
    if relevant_exp:
        exp_pts, exp_why = w["experience"], "Has related work or internship experience"
    elif experience:
        exp_pts, exp_why = w["experience"] / 2, "Has some experience, not related to this area"
    else:
        exp_pts, exp_why = 0, "No work experience (normal for a student)"

    parts = {
        "skills": {"points": round(skills_pts), "max": w["skills"],
                   "why": f"Found {len(found)} of {len(required)} required skills"
                          + (f" and {len(preferred_found)} preferred" if preferred_found else "")},
        "projects": {"points": round(projects_pts), "max": w["projects"],
                     "why": f"{len(relevant_projects)} project(s) related to this area"},
        "education": {"points": round(edu_pts), "max": w["education"], "why": edu_why},
        "experience": {"points": round(exp_pts), "max": w["experience"], "why": exp_why},
    }
    total = round(skills_pts + projects_pts + edu_pts + exp_pts)
    return total, parts, gaps


def shortlist_status(score: int) -> str:
    if score >= config.SHORTLIST_AT:
        return "shortlisted"
    if score >= config.BORDERLINE_AT:
        return "borderline"
    return "not_shortlisted_check"


def best_area_by_rules(profile: dict, cv_text: str, areas: list) -> tuple[str, int, float]:
    """The area with the highest score by rules. Returns (name, score, share of required skills found)."""
    best = ("", -1, 0.0)
    cv = CVText(cv_text)
    for a in areas:
        total, _, _ = score_cv(profile, cv_text, a)
        if total is None:
            continue
        req = a.required_skills or []
        share = sum(1 for s in req if cv.has(s)) / len(req)
        if total > best[1]:
            best = (a.name, total, share)
    return best


def area_fit_matrix(profile: dict, cv_text: str, areas: list) -> list[dict]:
    """Score the CV against every active area so the HR officer can compare them side by side."""
    if not cv_text or not cv_text.strip():
        return []
    cv = CVText(cv_text)
    rows: list[dict] = []
    for a in areas:
        total, _, gaps = score_cv(profile, cv_text, a)
        if total is None:
            continue
        req = list(a.required_skills or [])
        pref = list(a.preferred_skills or [])
        found_req = [s for s in req if cv.has(s)]
        found_pref = [s for s in pref if cv.has(s)]
        rows.append({
            "area": a.name,
            "score": total,
            "status": shortlist_status(total),
            "matched_required": found_req,
            "total_required": len(req),
            "matched_preferred": found_pref,
            "gaps": gaps[:4],
        })
    # Highest score first; break ties by required skills found, then preferred skills.
    rows.sort(key=lambda r: (r["score"], len(r["matched_required"]), len(r["matched_preferred"])), reverse=True)
    return rows


QUESTION_STOP = {"what", "which", "where", "when", "who", "how", "why", "does", "did", "do", "has", "have",
                 "had", "can", "could", "would", "should", "is", "are", "was", "were", "the", "this", "that",
                 "student", "candidate", "applicant", "about", "tell", "me", "us", "any", "their", "his", "her",
                 "show", "check", "list", "find", "mention", "mentioned", "cv", "resume"}


def rule_answer_cv(question: str, cv_text: str, profile: dict | None = None) -> dict:
    """Answer a question from the CV by plain rules when the model is offline or fails the checks."""
    p = profile or {}
    q_norm = norm(question)
    keywords = [w for w in q_norm.split() if len(w) >= 2 and w not in QUESTION_STOP]
    lines = [re.sub(r"^[\-•*●▪>]+\s*", "", ln).strip() for ln in (cv_text or "").splitlines() if len(ln.strip()) >= 8]
    cv = CVText(cv_text)

    # Broad summary / strengths questions use the parsed profile sections.
    if any(w in q_norm for w in ("summary", "summarize", "strength", "strengths", "overview", "highlight", "highlights")):
        skills = (p.get("skills") or [])[:6]
        edu = (p.get("education") or [])[:1]
        projs = (p.get("projects") or [])[:2]
        quotes = [x for x in (edu + projs + skills[:1]) if x and cv.has(x)][:3]
        parts = []
        if edu:
            parts.append(f"Education: {edu[0]}.")
        if skills:
            parts.append(f"Key skills listed in the CV include {', '.join(skills)}.")
        if projs:
            parts.append(f"Notable project work: {projs[0]}.")
        ans = " ".join(parts) if parts else "The CV lists the candidate's academic background, skills and projects."
        return {"answer": ans, "quotes": quotes}

    # Specific questions match CV lines that contain the question keywords.
    scored: list[tuple[int, str]] = []
    for ln in lines:
        ln_cv = CVText(ln)
        hits = sum(1 for k in keywords if ln_cv.has(k))
        if hits > 0:
            scored.append((hits, ln))
    scored.sort(key=lambda x: x[0], reverse=True)
    quotes = []
    for _, ln in scored:
        if ln not in quotes and cv.has(ln):
            quotes.append(ln[:220])
        if len(quotes) >= 3:
            break
    if quotes:
        return {
            "answer": f"Based on the CV, here is what is written regarding your question: {' '.join(quotes[:2])}",
            "quotes": quotes,
        }
    return {
        "answer": "This topic is not mentioned in the candidate's CV. Only facts explicitly written in the CV can be confirmed.",
        "quotes": [],
    }


# ---- Interview slots -------------------------------------------------------------------
def format_slot_label(iso: str) -> str:
    """Turn an ISO slot string into a friendly date and time label."""
    dt = datetime.fromisoformat(iso)
    return f"{dt.strftime('%A')} {dt.day} {dt.strftime('%B %Y')} at {dt.strftime('%H:%M')}"


def upcoming_interview_slots(taken: set[str], count: int = 21, today: date | None = None) -> list[dict[str, str]]:
    """Next free slots on working days, starting tomorrow. Returns [{'iso': ..., 'label': ...}]."""
    out: list[dict[str, str]] = []
    day = (today or date.today()) + timedelta(days=1)
    for _ in range(60):
        if day.weekday() < 5:
            for t in config.INTERVIEW_TIMES:
                iso = f"{day.isoformat()}T{t}"
                if iso not in taken:
                    out.append({"iso": iso, "label": format_slot_label(iso)})
                    if len(out) >= count:
                        return out
        day += timedelta(days=1)
    return out


def pick_interview_slot(taken: set[str], today: date | None = None) -> tuple[str, str]:
    """Next free slot on a working day, starting tomorrow. Returns (iso, label)."""
    slots = upcoming_interview_slots(taken, count=1, today=today)
    if slots:
        return slots[0]["iso"], slots[0]["label"]
    raise RuntimeError("No free interview slots in the next 60 days.")


# ---- Onboarding --------------------------------------------------------------------------
def induction_checklist(area) -> list[dict]:
    name = area.name if area else "the team"
    items = [
        "Send the internship agreement and get it signed",
        "Collect the university letter confirming the internship period",
        "Create the intern's SLT accounts and email",
        f"Assign a mentor in {name}",
        f"Add the intern to the {name} team channel",
        "Book a place in the next induction session",
    ]
    for extra in (area.checklist if area else []) or []:
        if extra not in items:
            items.append(extra)
    return [{"item": i, "done": False} for i in items]


# ---- Reply templates (used when the model's draft fails the checks) ----------------
SIGN_OFF = "Kind regards,\nSLT Internship Team"


def template_draft(kind: str, facts: dict) -> dict:
    first = (facts.get("name") or "there").split()[0]
    area = facts.get("area") or "the internship"
    app_no = facts.get("app_no", "")
    if kind == "invitation":
        return {"subject": f"Interview invitation - {area} internship",
                "body": (f"Dear {first},\n\nThank you for applying for an internship at SLT. We are pleased to "
                         f"invite you to an interview for the {area} area.\n\nDate and time: {facts.get('interview')}\n"
                         f"Application number: {app_no}\n\nPlease reply to this email to confirm that you can "
                         f"attend. If the time does not suit you, tell us and we will find another slot.\n\n{SIGN_OFF}")}
    if kind == "acceptance":
        return {"subject": f"Internship offer - {area}",
                "body": (f"Dear {first},\n\nCongratulations. We are happy to tell you that you have been selected "
                         f"for an internship in the {area} area at SLT.\n\nYour work ID is {facts.get('work_id')}. "
                         f"Please use it in all messages with us.\n\nOur team will contact you soon with the "
                         f"induction details and the documents to sign.\n\n{SIGN_OFF}")}
    if kind == "hold":
        return {"subject": "Update on your internship interview",
                "body": (f"Dear {first},\n\nThank you for attending the interview for the {area} internship. "
                         f"Your result is on hold for now, and we will contact you as soon as a decision is made. "
                         f"You do not need to do anything at this stage.\n\n{SIGN_OFF}")}
    if kind == "rejection":
        return {"subject": "Your internship interview result",
                "body": (f"Dear {first},\n\nThank you for attending the interview for the {area} internship and for "
                         f"your interest in SLT. After careful thought, we are not able to offer you a place this "
                         f"time. We wish you every success in your studies and future career.\n\n{SIGN_OFF}")}
    # not shortlisted
    return {"subject": "Your internship application",
            "body": (f"Dear {first},\n\nThank you for applying for an internship at SLT. We have reviewed your "
                     f"application carefully, and we are not able to invite you to an interview this time. We "
                     f"encourage you to keep building your skills and to apply again in the future.\n\n{SIGN_OFF}")}


def template_questions(area_name: str, gaps: list[str], profile: dict) -> list[str]:
    qs = [f"What made you interested in an internship in {area_name}?"]
    projects = profile.get("projects") or []
    if projects:
        qs.append(f"Tell us about your project \"{projects[0][:80]}\". What was your own part in it?")
    else:
        qs.append("Tell us about a project you are proud of. What was your own part in it?")
    for g in gaps[:2]:
        qs.append(f"How much do you know about {g}, and how would you learn more about it?")
    qs.append("Describe a problem you could not solve at first. What did you do?")
    qs.append("How do you plan your work when you have several deadlines?")
    return qs[:8]
