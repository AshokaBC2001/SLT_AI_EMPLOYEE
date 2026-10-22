"""Builds the sample CV files in samples/ from the texts in samples/src/.

Most samples stay as .txt. One is turned into a Word file and one into a PDF,
so all three file types are tested. All people in the samples are made up.
    python tests/make_samples.py
"""
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "samples" / "src"
OUT = ROOT / "samples"

AS_DOCX = {"dilini_wickramasinghe"}
AS_PDF = {"tharushi_madushani"}


def to_docx(text: str, path: Path) -> None:
    import docx
    d = docx.Document()
    lines = text.split("\n")
    d.add_heading(lines[0], level=1)
    for line in lines[1:]:
        d.add_paragraph(line)
    d.save(str(path))


def to_pdf(text: str, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    y = A4[1] - 60
    for i, line in enumerate(text.split("\n")):
        c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 14 if i == 0 else 10)
        while len(line) > 100:
            cut = line.rfind(" ", 0, 100)
            c.drawString(50, y, line[:cut])
            line = line[cut + 1:]
            y -= 14
        c.drawString(50, y, line)
        y -= 18 if i == 0 else 14
    c.save()


for src in sorted(SRC.glob("*.txt")):
    text = src.read_text(encoding="utf-8")
    if src.stem in AS_DOCX:
        to_docx(text, OUT / f"{src.stem}.docx")
    elif src.stem in AS_PDF:
        to_pdf(text, OUT / f"{src.stem}.pdf")
    else:
        (OUT / src.name).write_text(text, encoding="utf-8")
print("Samples written to", OUT)
